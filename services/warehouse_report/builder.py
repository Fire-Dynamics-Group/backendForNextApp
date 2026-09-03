"""Turn a rendered line-oriented text into paragraphs on a Word "skin".

Line grammar (one paragraph per line, blank lines ignored)::

    # Heading                    Main Headd            (numbered 1., 2., ...)
    ## Heading                   FD Subheading 1       (numbered 3.1., ...)
    ### Heading                  Astute References, unnumbered, bold, indented
    = text                       definition line (ASTUTE Body Text), indented
    > text                       equation "where:" line, unnumbered, italic, indented
    - text                       reference list entry, hanging indent
    @Style Name| text            any paragraph style from the skin, as-is
    !! text                      engineer prompt: bold, highlighted, must be edited out
    !fig key | caption           picture from figures[key], then "Figure N: caption"
    !figblock key | blockname    a block that is itself a figure with its own caption;
                                 {{FIGURE_NUMBER}} inside it becomes N
    !tab key                     marker: assigns the next table number to key
    !pagebreak                   start a new page
    !appendix A                  from here: figures/tables number A.1, A.2, ...; "##"
                                 headings number A.1. and body paragraphs A.1.1.
    !table key | caption         table from tables[key] (rows of cells; "^" merges up)
    [[name]]                     a block from blocks/: a whole paragraph or table when
                                 the fragment is w:p / w:tbl, a display equation when
                                 m:oMathPara, or an inline symbol (m:oMath) mid-line
    text                         Parag                 (numbered 1.1.1., ...)

Inline: ``**bold**``; ``[3]`` citations are superscript (except in reference lines);
``{fig:key}`` / ``{tab:key}`` become the figure/table numbers.
``{{KEY}}`` placeholders inside w:tbl / w:p blocks and in the skin's cover/header are
replaced from the substitution mapping.

The per-line paragraph overrides (no numbering, indents, italics) reproduce what the
team's template applied directly to those paragraphs on top of the named styles.
"""

import copy
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Optional, Tuple, Union

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips

FIGURE_WIDTH_IN = 6.0
PLACEHOLDER = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
INLINE = re.compile(r"(\[\[\w+\]\]|\*\*.+?\*\*|\[\d+\])")
CITATION = re.compile(r"^\[\d+\]$")

ImageSource = Union[BytesIO, str, Path]


@dataclass(frozen=True)
class Fmt:
    """Paragraph + run overrides applied on top of a named style."""

    style: str
    unnumbered: bool = False
    indent: Optional[int] = None  # twips
    hanging: Optional[int] = None  # twips
    center: bool = False
    bold: bool = False
    italic: bool = False
    highlight: bool = False
    font: Optional[str] = None
    size_pt: Optional[float] = None
    superscript_citations: bool = True  # "[3]" in running text; off for the reference list itself
    numbering: Optional[Tuple[int, int]] = None  # (ilvl, numId) applied directly, e.g. an appendix list


HEADING_1 = Fmt("Main Headd")
HEADING_2 = Fmt("FD Subheading 1")
HEADING_3 = Fmt("Astute References", unnumbered=True, indent=720, bold=True)
BODY = Fmt("Parag")
DEFINITION = Fmt("ASTUTE Body Text", indent=720, font="Segoe UI Semilight", size_pt=10)
WHERE_LINE = Fmt("Appendix Text", unnumbered=True, indent=1416, italic=True)
EQUATION = Fmt("Appendix Text", unnumbered=True, indent=720)
PICTURE = Fmt("Appendix Text", unnumbered=True, center=True)
CAPTION = Fmt("figure / table title")
REFERENCE = Fmt("List Paragraph", indent=714, hanging=357, superscript_citations=False)
ENGINEER_NOTE = Fmt("Parag", unnumbered=True, indent=720, bold=True, highlight=True)
TABLE_STYLE = "FD Table"  # the skin's own table style

# After "!appendix X": level-2 headings and body paragraphs keep their styles but are
# put on a clone of the main list whose level texts read "X.%2." and "X.%2.%3.", the
# way the team's appendix template numbers A.1. / A.1.1. (see _add_appendix_list).


class DocBuilder:
    def __init__(self, skin: Union[str, Path], blocks_dir: Union[str, Path]):
        self.doc = Document(str(skin))
        self.blocks_dir = Path(blocks_dir)
        self._body = self.doc.element.body
        self._sect_pr = self._body.find(qn("w:sectPr"))
        self._appendix_num: Optional[int] = None

    # ------------------------------------------------------------------ public

    def substitute(self, mapping: Mapping[str, str]) -> None:
        """Replace {{KEY}} placeholders in the skin: body, text boxes, headers, footers."""
        roots = [self.doc.element]
        for section in self.doc.sections:
            roots.extend([section.header._element, section.footer._element])
        for root in roots:
            for paragraph in root.iter(qn("w:p")):
                _substitute_in_paragraph(paragraph, mapping)

    def render(
        self,
        text: str,
        figures: Optional[Mapping[str, ImageSource]] = None,
        tables: Optional[Mapping[str, List[List[str]]]] = None,
        substitutions: Optional[Mapping[str, str]] = None,
    ) -> None:
        figures = figures or {}
        tables = tables or {}
        substitutions = substitutions or {}

        lines = [line.rstrip() for line in text.splitlines()]
        lines = _resolve_numbers(lines)

        for line in lines:
            if not line.strip():
                continue
            self._render_line(line, figures, tables, substitutions)

    def save(self) -> BytesIO:
        stream = BytesIO()
        self.doc.save(stream)
        stream.seek(0)
        return stream

    def text(self) -> str:
        """Plain text of the body, for tests and snapshots."""
        parts = []
        for child in self._body.iterchildren():
            if child.tag == qn("w:p"):
                parts.append(_paragraph_text(child))
            elif child.tag == qn("w:tbl"):
                for row in child.iter(qn("w:tr")):
                    parts.append(" | ".join(_paragraph_text(p) for p in row.iter(qn("w:p"))))
        return "\n".join(parts)

    # ----------------------------------------------------------------- private

    def _render_line(self, line, figures, tables, substitutions) -> None:
        if line.startswith("!fig "):
            key, caption = _split_marker(line[5:])
            self._add_picture(figures[key])
            self._add_paragraph(CAPTION, f"Figure {caption}")
        elif line.startswith("!figblock "):
            name, number = _split_marker(line[10:])
            self._add_standalone_block(name, {**substitutions, "FIGURE_NUMBER": number})
        elif line.startswith("!tab "):
            return  # numbering marker only
        elif line.strip() == "!pagebreak":
            self.doc.add_page_break()
        elif line.startswith("!appendix "):
            self._appendix_num = self._add_appendix_list(line[10:].strip())
        elif line.startswith("!table "):
            key, caption = _split_marker(line[7:])
            self._add_table(tables[key])
            self._add_paragraph(CAPTION, f"Table {caption}")
        elif line.startswith("!! "):
            self._add_paragraph(ENGINEER_NOTE, "ENGINEER TO EDIT: " + line[3:].strip())
        elif line.startswith("### "):
            self._add_paragraph(HEADING_3, line[4:])
        elif line.startswith("## "):
            fmt = HEADING_2 if self._appendix_num is None else Fmt(HEADING_2.style, numbering=(1, self._appendix_num))
            self._add_paragraph(fmt, line[3:])
        elif line.startswith("# "):
            self._add_paragraph(HEADING_1, line[2:])
        elif line.startswith("= "):
            self._add_paragraph(DEFINITION, line[2:])
        elif line.startswith("> "):
            self._add_paragraph(WHERE_LINE, line[2:])
        elif line.startswith("- "):
            self._add_paragraph(REFERENCE, line[2:])
        elif line.startswith("@"):
            # "@Style| text"; a trailing "*" on the style name suppresses its list number
            # (e.g. "@Main Headd*| References" for a top-level heading with no "5.").
            style, content = line[1:].split("|", 1)
            style = style.strip()
            unnumbered = style.endswith("*")
            self._add_paragraph(Fmt(style.rstrip("*").strip(), unnumbered=unnumbered), content.strip())
        elif re.fullmatch(r"\[\[\w+\]\]", line.strip()):
            self._add_standalone_block(line.strip()[2:-2], substitutions)
        else:
            fmt = BODY if self._appendix_num is None else Fmt(BODY.style, numbering=(2, self._appendix_num))
            self._add_paragraph(fmt, line)

    def _add_paragraph(self, fmt: Fmt, content: str):
        paragraph = self.doc.add_paragraph(style=fmt.style)
        _apply_paragraph_format(paragraph, fmt)
        for token in INLINE.split(content):
            if not token:
                continue
            if token.startswith("[[") and token.endswith("]]"):
                element = self._load_block(token[2:-2])
                if element.tag in (qn("m:oMath"), qn("m:oMathPara")):
                    paragraph._p.append(element)
                else:
                    raise ValueError(f"block {token} is not inline maths; put it on its own line")
            elif token.startswith("**") and token.endswith("**"):
                run = paragraph.add_run(token[2:-2])
                _apply_run_format(run, fmt)
                run.bold = True
            elif CITATION.match(token) and fmt.superscript_citations:
                run = paragraph.add_run(token)
                _apply_run_format(run, fmt)
                run.font.superscript = True
            else:
                _apply_run_format(paragraph.add_run(token), fmt)
        return paragraph

    def _add_standalone_block(self, name: str, substitutions: Mapping[str, str]) -> None:
        element = self._load_block(name)
        if element.tag == qn("m:oMathPara"):
            paragraph = self.doc.add_paragraph(style=EQUATION.style)
            _apply_paragraph_format(paragraph, EQUATION)
            paragraph._p.append(element)
        elif element.tag in (qn("w:p"), qn("w:tbl")):
            for paragraph in element.iter(qn("w:p")):
                _substitute_in_paragraph(paragraph, substitutions)
            if element.tag == qn("w:tbl"):
                _keep_table_together(element)
                self._keep_last_paragraph_with_next()
            self._sect_pr.addprevious(element)
        else:
            raise ValueError(f"block {name} ({element.tag}) cannot stand alone")

    def _add_picture(self, source: ImageSource) -> None:
        paragraph = self.doc.add_paragraph(style=PICTURE.style)
        _apply_paragraph_format(paragraph, PICTURE)
        paragraph.paragraph_format.keep_with_next = True  # caption stays under its figure
        if isinstance(source, BytesIO):
            source.seek(0)
        else:
            source = str(source)
        paragraph.add_run().add_picture(source, width=Inches(FIGURE_WIDTH_IN))

    def _add_appendix_list(self, letter: str) -> int:
        """Clone the main heading/paragraph list so its levels read "A.1." and "A.1.1.".

        Word numbers by list instance, so the body's 3.5.3 and the appendix's A.5.3
        cannot share one. The clone keeps the main list's indents and fonts, drops the
        style links (or every Parag in the body would join it), restarts its counters
        and prefixes the level text with the appendix letter.
        """
        numbering = self.doc.part.numbering_part.element

        def num_for(style_name: str):
            num_id = self.doc.styles[style_name].element.find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
            return next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id)

        # Level 1 (headings) from the heading style's list and level 2 (paragraphs)
        # from the body style's: the two body styles number through different list
        # instances of the same abstract list, with different number fonts/colours.
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        clone = copy.deepcopy(num_for(BODY.style))
        clone.set(qn("w:numId"), str(new_id))
        heading_level = next(o for o in num_for(HEADING_2.style).findall(qn("w:lvlOverride")) if o.get(qn("w:ilvl")) == "1")
        for override in clone.findall(qn("w:lvlOverride")):
            if override.get(qn("w:ilvl")) == "1":
                override.addprevious(copy.deepcopy(heading_level))
                clone.remove(override)
        for override in clone.findall(qn("w:lvlOverride")):
            level = int(override.get(qn("w:ilvl")))
            lvl = override.find(qn("w:lvl"))
            if lvl is None:
                continue
            style = lvl.find(qn("w:pStyle"))
            if style is not None:
                lvl.remove(style)
            if level in (1, 2):
                lvl.find(qn("w:lvlText")).set(qn("w:val"), f"{letter}.%2." if level == 1 else f"{letter}.%2.%3.")
                restart = OxmlElement("w:startOverride")
                restart.set(qn("w:val"), "1")
                override.insert(0, restart)
        cleanup = numbering.find(qn("w:numIdMacAtCleanup"))
        if cleanup is not None:
            cleanup.addprevious(clone)
        else:
            numbering.append(clone)
        return new_id

    def _keep_last_paragraph_with_next(self) -> None:
        """The sentence introducing a table ("...are provided in Table 1.") moves with it."""
        previous = self._sect_pr.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            Paragraph(previous, self.doc._body).paragraph_format.keep_with_next = True

    def _add_table(self, rows: List[List[str]]) -> None:
        self._keep_last_paragraph_with_next()
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = TABLE_STYLE
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if value == "^":
                    continue
                table.cell(r, c).text = value
                if r == 0:
                    for run in table.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        # Merge "^" cells upward into the nearest filled cell above.
        for c in range(len(rows[0])):
            top = 0
            for r in range(len(rows)):
                if rows[r][c] != "^":
                    top = r
                elif r == len(rows) - 1 or rows[r + 1][c] != "^":
                    table.cell(top, c).merge(table.cell(r, c))
        _keep_table_together(table._tbl)

    def _load_block(self, name: str):
        path = self.blocks_dir / f"{name}.xml"
        if not path.exists():
            raise KeyError(f"unknown block [[{name}]] (no {path.name})")
        return parse_xml(path.read_text(encoding="utf-8"))


# ------------------------------------------------------------------ helpers


def _apply_paragraph_format(paragraph, fmt: Fmt) -> None:
    pf = paragraph.paragraph_format
    numbering = (0, 0) if fmt.unnumbered else fmt.numbering
    if numbering is not None:
        ppr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(numbering[0]))
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), str(numbering[1]))
        num_pr.append(ilvl)
        num_pr.append(num_id)
        ppr.insert_element_before(
            num_pr,
            "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
            "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
            "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
            "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
            "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
        )
    if fmt.indent is not None:
        pf.left_indent = Twips(fmt.indent)
    if fmt.hanging is not None:
        pf.first_line_indent = Twips(-fmt.hanging)
    if fmt.center:
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _keep_table_together(tbl) -> None:
    """Keep a table on one page and its caption with it.

    Word has no table-level setting for this: rows get cantSplit, and every
    paragraph in every row gets keepNext, so the whole table (and the caption
    paragraph after it) moves to the next page rather than breaking mid-way.
    """
    for row in tbl.findall(qn("w:tr")):
        tr_pr = row.find(qn("w:trPr"))
        if tr_pr is None:
            tr_pr = OxmlElement("w:trPr")
            tbl_pr_ex = row.find(qn("w:tblPrEx"))
            if tbl_pr_ex is not None:
                tbl_pr_ex.addnext(tr_pr)
            else:
                row.insert(0, tr_pr)
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.insert(0, OxmlElement("w:cantSplit"))
        for paragraph in row.iter(qn("w:p")):
            p_pr = paragraph.find(qn("w:pPr"))
            if p_pr is None:
                p_pr = OxmlElement("w:pPr")
                paragraph.insert(0, p_pr)
            if p_pr.find(qn("w:keepNext")) is None:
                keep = OxmlElement("w:keepNext")
                style = p_pr.find(qn("w:pStyle"))
                if style is not None:
                    style.addnext(keep)
                else:
                    p_pr.insert(0, keep)


def _apply_run_format(run, fmt: Fmt) -> None:
    if fmt.bold:
        run.bold = True
    if fmt.italic:
        run.italic = True
    if fmt.highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if fmt.font:
        run.font.name = fmt.font
    if fmt.size_pt:
        run.font.size = Pt(fmt.size_pt)


def _split_marker(rest: str):
    key, _, caption = rest.partition("|")
    return key.strip(), caption.strip()


def _resolve_numbers(lines: Iterable[str]) -> List[str]:
    """Assign figure/table numbers in order of appearance and resolve references."""
    lines = list(lines)
    figures: Dict[str, str] = {}
    tables: Dict[str, str] = {}
    prefix = ""
    fig_count = tab_count = 0
    for line in lines:
        if line.startswith("!appendix "):
            # "!appendix A": numbering restarts as A.1, A.2, ... from here on.
            prefix = line[10:].strip() + "."
            fig_count = tab_count = 0
        elif line.startswith("!fig ") or line.startswith("!figblock "):
            key, _ = _split_marker(line.split(" ", 1)[1])
            if key not in figures:
                fig_count += 1
                figures[key] = f"{prefix}{fig_count}"
        elif line.startswith("!tab ") or line.startswith("!table "):
            # "!tab key" reserves the next number ahead of a block table; "!table key |
            # caption" takes it when it is the first mention.
            key = line[5:].strip() if line.startswith("!tab ") else _split_marker(line[7:])[0]
            if key not in tables:
                tab_count += 1
                tables[key] = f"{prefix}{tab_count}"

    def resolve(text: str) -> str:
        text = re.sub(r"\{fig:(\w+)\}", lambda m: figures[m.group(1)], text)
        return re.sub(r"\{tab:(\w+)\}", lambda m: tables[m.group(1)], text)

    out = []
    for line in lines:
        if line.startswith("!fig "):
            key, caption = _split_marker(line[5:])
            out.append(f"!fig {key} | {figures[key]}: {resolve(caption)}")
        elif line.startswith("!figblock "):
            key, name = _split_marker(line[10:])
            out.append(f"!figblock {name} | {figures[key]}")
        elif line.startswith("!table "):
            key, caption = _split_marker(line[7:])
            out.append(f"!table {key} | {tables[key]}: {resolve(caption)}")
        else:
            out.append(resolve(line))
    return out


def _direct_texts(paragraph) -> list:
    """The w:t nodes of this paragraph's own runs.

    Deliberately not ``paragraph.iter(w:t)``: a paragraph that anchors text boxes
    (the cover page, the ASET/RSET diagram) contains whole nested paragraphs, and
    those must be handled on their own, not folded into their anchor's text.
    """
    p_tag = qn("w:p")

    def owner(node):
        parent = node.getparent()
        while parent is not None and parent.tag != p_tag:
            parent = parent.getparent()
        return parent

    return [t for t in paragraph.iter(qn("w:t")) if owner(t) is paragraph]


def _paragraph_text(paragraph) -> str:
    return "".join(t.text or "" for t in paragraph.iter(qn("w:t")))


def _substitute_in_paragraph(paragraph, mapping: Mapping[str, str]) -> None:
    """Replace {{KEY}} even when Word has split it across runs.

    The paragraph's own text is joined, substituted, and written back into the
    first text node; the other text nodes are emptied. Run formatting of the first
    node is kept, which is what the cover page and table cells need.
    """
    texts = _direct_texts(paragraph)
    if not texts:
        return
    joined = "".join(t.text or "" for t in texts)
    if "{{" not in joined:
        return

    def replacement(match):
        key = match.group(1)
        return str(mapping[key]) if key in mapping else match.group(0)

    new = PLACEHOLDER.sub(replacement, joined)
    if new == joined:
        return
    texts[0].text = new
    texts[0].set(qn("xml:space"), "preserve")
    for t in texts[1:]:
        t.text = ""
