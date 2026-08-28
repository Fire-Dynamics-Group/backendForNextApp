"""Render the warehouse smoke layer Word report.

This module formats numbers it is given — it does not run the zone model. The
calculation lives in the browser (fd-toolstation, lib/smoke-layer-calc.ts) and the
results arrive with the request, so the report always matches what the engineer saw.
Charts here are re-plotted with matplotlib purely to get print-quality figures.
"""

from datetime import date
from io import BytesIO
from typing import List, Optional, Tuple

import matplotlib

matplotlib.use("Agg")  # no display on Railway
import matplotlib.pyplot as plt  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

from models.smoke_layer_models import (  # noqa: E402
    SmokeLayerInputs,
    SmokeLayerResults,
    SmokeLayerStep,
)

# Categorical slots 1 and 2 of the toolstation chart palette, so the report and
# the on-screen charts read as the same set of figures.
SERIES_1 = "#2a78d6"
SERIES_2 = "#eb6834"
ANNOTATION = "#8a8781"

TENABILITY_HEIGHT = 2.0
FIGURE_WIDTH_IN = 6.0


def _plot(
    steps: List[SmokeLayerStep],
    series: List[Tuple[str, str, str]],
    ylabel: str,
    thresholds: Optional[List[Tuple[float, str]]] = None,
    markers: Optional[List[Tuple[float, str]]] = None,
    from_zero: bool = False,
) -> BytesIO:
    """Plot one figure and return it as a PNG stream.

    `series` is a list of (attribute, label, colour).
    """
    times = [s.time for s in steps]

    fig, ax = plt.subplots(figsize=(FIGURE_WIDTH_IN, 3.0), dpi=200)

    for attr, label, colour in series:
        ax.plot(times, [getattr(s, attr) for s in steps], color=colour, linewidth=1.6, label=label)

    for value, label in thresholds or []:
        ax.axhline(value, color=ANNOTATION, linewidth=1, linestyle="--")
        ax.annotate(
            label,
            xy=(times[0], value),
            xytext=(3, 3),
            textcoords="offset points",
            fontsize=7,
            color=ANNOTATION,
        )

    for value, label in markers or []:
        ax.axvline(value, color=ANNOTATION, linewidth=1, linestyle="--")
        ax.annotate(
            label,
            xy=(value, ax.get_ylim()[1]),
            xytext=(-3, -10),
            textcoords="offset points",
            fontsize=7,
            color=ANNOTATION,
            ha="right",
        )

    ax.set_xlabel("Time (s)", fontsize=8)
    ax.set_ylabel(ylabel, fontsize=8)
    ax.tick_params(labelsize=7)
    ax.grid(True, color="#eceae5", linewidth=0.8)
    ax.set_axisbelow(True)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    for spine in ("left", "bottom"):
        ax.spines[spine].set_color("#d6d3cc")
    if from_zero:
        ax.set_ylim(bottom=0)
    if len(series) > 1:
        ax.legend(fontsize=7, frameon=False)

    fig.tight_layout()
    stream = BytesIO()
    fig.savefig(stream, format="png")
    plt.close(fig)
    stream.seek(0)
    return stream


def _add_table(document: Document, rows: List[Tuple[str, str]], heading: str) -> None:
    document.add_heading(heading, level=2)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for paragraph in cells[0].paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)
        for paragraph in cells[1].paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)


def _seconds(value: float) -> str:
    return f"{round(value):,} s"


def generate_smoke_layer_report(
    project_name: str,
    engineer_name: str,
    inputs: SmokeLayerInputs,
    results: SmokeLayerResults,
) -> BytesIO:
    """Build the Word report and return it as a BytesIO stream."""

    document = Document()

    document.add_heading("Warehouse Smoke Layer Assessment", level=0)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle_run = subtitle.add_run(
        f"{project_name or 'Untitled project'}\n"
        f"{engineer_name or 'Fire Dynamics Group'} — {date.today().strftime('%d %B %Y')}"
    )
    subtitle_run.font.size = Pt(10)

    # ---- Methodology ----
    document.add_heading("Methodology", level=2)
    document.add_paragraph(
        "The height of the smoke layer has been estimated using a single-zone model "
        "following the methodology of Drysdale, Introduction to Fire Dynamics (2011). "
        "A t-squared design fire drives a plume whose mass flow is a function of the "
        "clear height below the layer. The smoke layer is treated as well mixed, its "
        "temperature being the mass-weighted average of the existing layer and the smoke "
        "entering it, and the plume temperature is capped at 900 K above ambient."
    )
    document.add_paragraph(
        "The depth of the layer is recalculated at each timestep from the whole "
        "accumulated smoke mass at the new layer temperature, rather than adding new "
        "smoke beneath a static older layer. The existing smoke therefore expands as the "
        "layer heats, driving the layer down faster. This is the more conservative "
        "assumption and is consistent with treating the layer as well mixed."
    )
    document.add_paragraph(
        "Racking is taken to displace volume uniformly over the full height of the "
        "compartment, and so reduces the plan area available to the smoke."
    )

    # ---- Inputs ----
    smoke_area = inputs.room_area * (1 - inputs.racking_perc)
    _add_table(
        document,
        [
            ("Floor area", f"{inputs.room_area:,.0f} m²"),
            ("Racking", f"{inputs.racking_perc * 100:.0f}% of volume"),
            ("Plan area available to smoke", f"{smoke_area:,.0f} m²"),
            ("Room height", f"{inputs.room_height:,.2f} m"),
            ("Fire growth rate", f"{inputs.fgr} kW/s²"),
            ("Convective fraction", "0.7"),
            ("Ambient temperature", "20 °C (293 K)"),
        ],
        "Room and fire",
    )

    _add_table(
        document,
        [
            ("Detection time", _seconds(inputs.detection_time)),
            ("Pre-movement time", _seconds(inputs.pre_movement_time)),
            ("Maximum travel distance", f"{inputs.maximum_travel_distance:,.1f} m"),
            ("Walking speed", f"{inputs.walking_speed} m/s"),
            ("Total exit width", f"{inputs.total_exit_width:,.2f} m"),
            ("Flow rate", f"{inputs.flow_rate} people/s/m"),
            ("Occupancy", f"{inputs.occupancy:,.0f} people"),
        ],
        "Means of escape",
    )

    # ---- Results ----
    document.add_heading("Results", level=2)

    verdict = (
        "The available safe escape time exceeds the required safe escape time."
        if results.margin_of_safety > 0
        else "The required safe escape time exceeds the available safe escape time. "
        "Tenability limits are exceeded before escape is complete."
    )
    document.add_paragraph(verdict)

    if results.aset_triggered:
        aset_text = _seconds(results.aset)
        aset_note = f"Smoke layer reaches {TENABILITY_HEIGHT:.0f} m"
    else:
        aset_text = f"> {_seconds(inputs.assessment_time)}"
        aset_note = (
            f"{TENABILITY_HEIGHT:.0f} m not reached within the assessment period; "
            f"the layer settles at {results.final_clear_height:.2f} m"
        )

    breach = (
        _seconds(results.breach_time)
        if results.reference_height_breached and results.breach_time is not None
        else "Not breached"
    )

    _add_table(
        document,
        [
            ("ASET", aset_text),
            ("", aset_note),
            ("RSET", _seconds(results.rset)),
            ("  Pre-evacuation time", _seconds(results.total_pre_evac)),
            ("  Queuing time", _seconds(results.queue_time)),
            ("  Occupant flow capacity", f"{results.people_per_second:,.2f} people/s"),
            (
                "Margin of safety",
                ("at least " if not results.aset_triggered else "")
                + f"{round(results.margin_of_safety):,} s",
            ),
            (f"Reference height ({inputs.reference_height:g} m)", breach),
            ("Clear height at end of run", f"{results.final_clear_height:.2f} m"),
        ],
        "ASET / RSET summary",
    )

    # ---- Figures ----
    document.add_heading("Figures", level=2)

    thresholds = [(TENABILITY_HEIGHT, f"{TENABILITY_HEIGHT:.0f} m tenability")]
    if inputs.reference_height != TENABILITY_HEIGHT:
        thresholds.append((inputs.reference_height, f"{inputs.reference_height:g} m reference"))

    markers = []
    if results.aset_triggered:
        markers.append((results.aset, f"ASET {round(results.aset):,} s"))
    if results.steps and results.rset <= results.steps[-1].time:
        markers.append((results.rset, f"RSET {round(results.rset):,} s"))

    figures = [
        (
            "Smoke layer height",
            _plot(
                results.steps,
                [("clear_height", "Clear height", SERIES_1)],
                "Clear height (m)",
                thresholds=thresholds,
                markers=markers,
                from_zero=True,
            ),
        ),
        (
            "Smoke layer temperature",
            _plot(
                results.steps,
                [
                    ("smoke_layer_temp", "Layer (well mixed)", SERIES_1),
                    ("added_smoke_temp", "Smoke entering layer", SERIES_2),
                ],
                "Temperature (°C)",
                markers=markers,
            ),
        ),
        (
            "Heat release rate",
            _plot(
                results.steps,
                [("hrr", "Total", SERIES_1), ("convective_hrr", "Convective", SERIES_2)],
                "Heat release rate (MW)",
                from_zero=True,
            ),
        ),
        (
            "Occupants escaped",
            _plot(
                results.steps,
                [("escaped", "Escaped", SERIES_1)],
                "People",
                markers=markers,
                from_zero=True,
            ),
        ),
    ]

    for index, (caption, stream) in enumerate(figures, start=1):
        document.add_picture(stream, width=Inches(FIGURE_WIDTH_IN))
        para = document.add_paragraph(f"Figure {index}: {caption}")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            run.font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
