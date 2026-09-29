"""End-to-end coverage for editing wording used by Word add-in proposals."""
import io

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient

from models.fee_proposal_models import (
    ClientDetails,
    CountryEnum,
    FeeOptions,
    FeeProposalRequest,
    ProjectDetails,
    ServiceConfig,
    DesignStagesRiba1to4,
)


@pytest.mark.asyncio
async def test_structural_services_have_separate_editable_word_text(monkeypatch, tmp_path):
    import database
    from main import app
    from routers import word_fee

    monkeypatch.setattr(database, "async_session", None)
    monkeypatch.setattr(word_fee, "LOCAL_WORD_STORE", tmp_path / "structural-word-text.json")

    prefixes = {
        "TMA_STRUCTURAL": "TMA Structural Fire Engineering",
        "TIME_EQUIVALENCY_STRUCTURAL": "Time-Equivalency Structural Fire Engineering",
        "FEM_STRUCTURAL": "FEM Structural Fire Engineering",
    }
    suffixes = (
        "PROPOSAL_INTRO",
        "FE_INTRO",
        "FE_SCOPE",
        "FE_SUB_BULLETS_1",
        "FE_SCOPE_2",
        "FE_SUB_BULLETS_2",
        "FE_SCOPE_3",
    )
    expected_groups = {
        f"{prefix}_{suffix}": group
        for prefix, group in prefixes.items()
        for suffix in suffixes
    }

    request = FeeProposalRequest(
        client=ClientDetails(first_name="Test", surname="Client", address_lines=["1 Test St"]),
        project=ProjectDetails(
            project_name="Test Tower",
            project_location="London",
            country=CountryEnum.ENGLAND_WALES,
        ),
        fee_options=FeeOptions(engineer_name="Sam Bennett", include_hourly_rates=False),
        design_stages_1_4=DesignStagesRiba1to4(
            tma_structural=ServiceConfig(included=True, fee=1000),
            time_equivalency_structural=ServiceConfig(included=True, fee=2000),
            fem_structural=ServiceConfig(included=True, fee=3000),
        ),
    )

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        listed = await client.get("/word/fee-proposal/text-blocks")
        assert listed.status_code == 200
        groups = {block["key"]: block["group_name"] for block in listed.json()}
        assert {key: groups.get(key) for key in expected_groups} == expected_groups

        for prefix in prefixes:
            saved = await client.put(
                f"/word/fee-proposal/text-blocks/{prefix}_PROPOSAL_INTRO",
                json={"content": f"{prefix} editable introduction", "edited_by": "Automated test"},
            )
            assert saved.status_code == 200
            saved = await client.put(
                f"/word/fee-proposal/text-blocks/{prefix}_FE_SCOPE",
                json={"content": f"{prefix} editable scope", "edited_by": "Automated test"},
            )
            assert saved.status_code == 200

        rendered = await client.post("/word/fee-proposal/render", json=request.model_dump(mode="json"))
        assert rendered.status_code == 200
        text = "\n".join(p.text for p in Document(io.BytesIO(rendered.content)).paragraphs)
        for prefix in prefixes:
            assert f"{prefix} editable introduction" in text
            assert f"{prefix} editable scope" in text

        web_rendered = await client.post("/fee-proposals/generate", json=request.model_dump(mode="json"))
        assert web_rendered.status_code == 200
        web_text = "\n".join(p.text for p in Document(io.BytesIO(web_rendered.content)).paragraphs)
        for prefix in prefixes:
            assert f"{prefix} editable introduction" not in web_text
            assert f"{prefix} editable scope" not in web_text


@pytest.mark.asyncio
async def test_word_text_edit_changes_addin_doc_only_and_can_be_reset(monkeypatch, tmp_path):
    import database
    from main import app
    from routers import word_fee

    monkeypatch.setattr(database, "async_session", None)
    monkeypatch.setattr(word_fee, "LOCAL_WORD_STORE", tmp_path / "word-text.json")

    request = FeeProposalRequest(
        client=ClientDetails(first_name="Test", surname="Client", address_lines=["1 Test St"]),
        project=ProjectDetails(
            project_name="Test Tower",
            project_location="London",
            country=CountryEnum.ENGLAND_WALES,
        ),
        fee_options=FeeOptions(engineer_name="Sam Bennett", include_hourly_rates=False),
        design_stages_1_4=DesignStagesRiba1to4(
            stage_1=ServiceConfig(included=True, fee=5000),
        ),
    )
    payload = request.model_dump(mode="json")
    edited_text = "WORD ADD-IN EDIT TEST SENTINEL"
    original_text = "I trust this provides you with the information that you require, however, should you wish to discuss, please do not hesitate to contact me."

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        listed = await client.get("/word/fee-proposal/text-blocks")
        assert listed.status_code == 200
        closing = next(block for block in listed.json() if block["key"] == "TERMS_CLOSING")

        saved = await client.put(
            "/word/fee-proposal/text-blocks/TERMS_CLOSING",
            json={"content": edited_text, "edited_by": "Automated test"},
        )
        assert saved.status_code == 200

        addin_doc = await client.post("/word/fee-proposal/render", json=payload)
        assert addin_doc.status_code == 200
        addin_text = "\n".join(p.text for p in Document(io.BytesIO(addin_doc.content)).paragraphs)
        assert edited_text in addin_text

        web_doc = await client.post("/fee-proposals/generate", json=payload)
        assert web_doc.status_code == 200
        web_text = "\n".join(p.text for p in Document(io.BytesIO(web_doc.content)).paragraphs)
        assert original_text in web_text
        assert edited_text not in web_text

        history = await client.get("/word/fee-proposal/text-blocks/TERMS_CLOSING/history")
        assert history.status_code == 200
        assert history.json()[0]["content"] == edited_text

        reset = await client.post(
            "/word/fee-proposal/text-blocks/TERMS_CLOSING/reset",
            json={"edited_by": "Automated test"},
        )
        assert reset.status_code == 200
        assert reset.json()["content"] == closing["content"]

        restored_doc = await client.post("/word/fee-proposal/render", json=payload)
        restored_text = "\n".join(p.text for p in Document(io.BytesIO(restored_doc.content)).paragraphs)
        assert original_text in restored_text
        assert edited_text not in restored_text
